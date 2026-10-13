"""
Exportación del procesador a Excel (XLSX) y CSV.

Excel: dos sheets — `CFDIs` (una fila por comprobante) y `Conceptos` (una fila
por concepto, ligada a UUID padre). Usa openpyxl en modo `write_only` para
streaming con bajo uso de memoria.

CSV: una fila por CFDI, con BOM UTF-8 para compatibilidad con Excel para
Windows. Opcionalmente expande conceptos.

`plantilla_a_xlsx` / `plantilla_a_csv`: el export por plantilla
(procesador/plantillas.py) — una hoja, columnas a elegir del usuario.

Los textos vienen del XML (los escribe el emisor): ninguno se escribe como
fórmula de Excel.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .db import CfdiFiltros, ProcesadorDB
from .plantillas import MAX_TEXTO_EXCEL, Columna, texto


# Tokens de la marca Señal (todoconta-apps/DESIGN.md), alineados con
# ui/src/app/globals.css: encabezados en tinta con texto blanco, como los
# botones de la app. El cian no va aquí (es solo para lo automático).
#
# Tipografía: la marca usa Geist, pero Excel no la trae ni se puede incrustar
# en un XLSX; Calibri es la sans-serif universal de Office (Windows y macOS)
# y la más cercana disponible sin instalar nada.
#
# Colores: hex literal (Excel los renderiza tal cual, mismo formato que CSS).
_BRAND_PRIMARY = "10141B"            # ink (tinta): encabezados
_BRAND_PRIMARY_FOREGROUND = "FFFFFF"  # texto sobre tinta
_BRAND_FONT_NAME = "Calibri"          # universal; fallback más cercano a Geist
_BRAND_FONT_SIZE = 11


# Columnas del sheet "CFDIs" (y del CSV plano)
_COLS_CFDI = [
    ("uuid", "UUID"),
    ("tipo", "Tipo"),
    ("fecha", "Fecha"),
    ("serie", "Serie"),
    ("folio", "Folio"),
    ("emisor_rfc", "RFC Emisor"),
    ("emisor_nombre", "Nombre Emisor"),
    ("receptor_rfc", "RFC Receptor"),
    ("receptor_nombre", "Nombre Receptor"),
    ("sub_total", "Subtotal"),
    ("descuento", "Descuento"),
    ("iva_trasladado", "IVA Trasladado"),
    ("ieps_trasladado", "IEPS Trasladado"),
    ("iva_retenido", "IVA Retenido"),
    ("isr_retenido", "ISR Retenido"),
    ("total", "Total"),
    ("moneda", "Moneda"),
    ("tipo_cambio", "Tipo de Cambio"),
    ("forma_pago", "Forma de Pago"),
    ("metodo_pago", "Método de Pago"),
    ("receptor_uso_cfdi", "Uso CFDI"),
    ("lugar_expedicion", "Lugar Expedición"),
    ("estado_sat", "Estado SAT"),
    # Flags de Comprobantes (derivadas en _enriquecer). Al FINAL para no
    # mover las posiciones de columnas en hojas de usuarios existentes.
    ("deducible", "Deducible"),
    ("estado_diot", "DIOT"),
]


def _enriquecer(cfdi: dict) -> dict:
    """Deriva las columnas display de los flags DIOT/deducibilidad (el dict
    ya trae elegible_diot/incluir_diot/deducible desde _row_to_dict)."""
    out = dict(cfdi)
    out["deducible"] = cfdi.get("deducible") or "Sin analizar"
    if not cfdi.get("elegible_diot"):
        out["estado_diot"] = "No aplica"
    else:
        out["estado_diot"] = "Pasa a DIOT" if cfdi.get("incluir_diot") else "Excluido"
    return out


_COLS_CONCEPTO = [
    ("cfdi_uuid", "UUID CFDI"),
    ("clave_prod_serv", "Clave Prod/Serv"),
    ("descripcion", "Descripción"),
    ("cantidad", "Cantidad"),
    ("clave_unidad", "Clave Unidad"),
    ("unidad", "Unidad"),
    ("valor_unitario", "Valor Unitario"),
    ("importe", "Importe"),
    ("descuento", "Descuento"),
]


def to_xlsx(
    db: ProcesadorDB,
    filtros: Optional[CfdiFiltros] = None,
) -> bytes:
    """
    Genera un XLSX con dos sheets (`CFDIs` y `Conceptos`) y devuelve los bytes.
    Modo write_only para bajo uso de memoria con +1000 CFDIs.
    """
    wb = Workbook(write_only=True)
    header_font = Font(
        bold=True,
        color=_BRAND_PRIMARY_FOREGROUND,
        name=_BRAND_FONT_NAME,
        size=_BRAND_FONT_SIZE,
    )
    header_fill = PatternFill("solid", fgColor=_BRAND_PRIMARY)

    # Sheet CFDIs
    ws_cfdis = wb.create_sheet("CFDIs")
    ws_cfdis.freeze_panes = "A2"
    headers = [label for _, label in _COLS_CFDI]
    ws_cfdis.append(_header_row(ws_cfdis, headers, font=header_font, fill=header_fill))

    uuids_cargados: list[str] = []
    for cfdi in map(_enriquecer, db.iter_all(filtros)):
        ws_cfdis.append([_celda(ws_cfdis, cfdi.get(key)) for key, _ in _COLS_CFDI])
        uuids_cargados.append(cfdi["uuid"])

    # Sheet Conceptos — cargado a partir de los UUIDs ya filtrados
    ws_conceptos = wb.create_sheet("Conceptos")
    ws_conceptos.freeze_panes = "A2"
    headers_c = [label for _, label in _COLS_CONCEPTO]
    ws_conceptos.append(_header_row(ws_conceptos, headers_c, font=header_font, fill=header_fill))

    if uuids_cargados:
        # El mismo uuid puede vivir bajo dos empresas — acotar al dueño
        # para no duplicar conceptos de la otra empresa en el export.
        mi_rfc = (filtros or {}).get("mi_rfc") or ""
        with db.cursor() as cur:
            # Chunks por límite SQLite de placeholders (999 por default)
            for chunk in _chunks(uuids_cargados, 500):
                placeholders = ",".join("?" for _ in chunk)
                cur.execute(
                    f"""
                    SELECT cfdi_uuid, clave_prod_serv, descripcion, cantidad,
                           clave_unidad, unidad, valor_unitario, importe, descuento
                    FROM conceptos
                    WHERE mi_rfc = ? AND cfdi_uuid IN ({placeholders})
                    """,
                    (mi_rfc, *chunk),
                )
                for row in cur.fetchall():
                    ws_conceptos.append([_celda(ws_conceptos, row[k]) for k, _ in _COLS_CONCEPTO])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def to_csv(
    db: ProcesadorDB,
    filtros: Optional[CfdiFiltros] = None,
    expandir_conceptos: bool = False,
) -> bytes:
    """
    Genera CSV con BOM UTF-8 (Excel-compat). Si `expandir_conceptos` es True,
    se emite una fila por concepto incluyendo las columnas del CFDI padre.
    """
    buf = io.StringIO()
    buf.write("﻿")  # BOM
    writer = csv.writer(buf)

    if not expandir_conceptos:
        writer.writerow([label for _, label in _COLS_CFDI])
        for cfdi in map(_enriquecer, db.iter_all(filtros)):
            writer.writerow([_csv_seguro(cfdi.get(key, "")) for key, _ in _COLS_CFDI])
    else:
        headers = (
            [label for _, label in _COLS_CFDI]
            + [label for k, label in _COLS_CONCEPTO if k != "cfdi_uuid"]
        )
        writer.writerow(headers)
        for cfdi in map(_enriquecer, db.iter_all(filtros)):
            base = [_csv_seguro(cfdi.get(key, "")) for key, _ in _COLS_CFDI]
            conceptos = db.conceptos_de(cfdi["uuid"], cfdi["mi_rfc"])
            if not conceptos:
                writer.writerow(base + [""] * (len(_COLS_CONCEPTO) - 1))
            else:
                for c in conceptos:
                    writer.writerow(
                        base + [_csv_seguro(c.get(key, ""))
                                for key, _ in _COLS_CONCEPTO if key != "cfdi_uuid"]
                    )

    return buf.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# Export por plantilla
# ---------------------------------------------------------------------------

_FORMATO_XLSX = {"fecha": "dd/mm/yyyy", "importe": "#,##0.00", "tipo_cambio": "#,##0.0000"}
_ANCHO_TIPO = {"fecha": 12, "importe": 14, "tipo_cambio": 11, "numero": 10}
_ANCHO_CAMPO = {
    "uuid": 38, "conceptos": 50, "descripcion": 45, "emisor_nombre": 32,
    "receptor_nombre": 32, "validacion_local": 30, "docs_relacionados": 40,
    "cfdi_relacionados": 40, "forma_pago": 26, "metodo_pago": 26, "pago_forma": 26,
    "uso_cfdi": 28, "tipo_comprobante": 16, "emisor_regimen": 30,
    "receptor_regimen": 30, "objeto_imp": 24,
}


def plantilla_a_xlsx(columnas: list[Columna], renglones: Iterable[list[Any]]) -> bytes:
    """Una hoja `CFDIs`: encabezado de marca congelado, autofiltro, fechas e
    importes como valores de Excel (no texto)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("CFDIs")
    for i, col in enumerate(columnas, 1):
        ancho = _ANCHO_CAMPO.get(col.campo.id) or _ANCHO_TIPO.get(col.campo.tipo, 16)
        ws.column_dimensions[get_column_letter(i)].width = max(
            ancho, min(len(col.encabezado) + 2, 40)
        )
    ws.freeze_panes = "A2"
    font, fill = _estilo_encabezado()
    ws.append(_header_row(ws, [c.encabezado for c in columnas], font=font, fill=fill))
    celdas = [_CeldasColumna(ws, col.campo.tipo) for col in columnas]
    total = 0
    for valores in renglones:
        ws.append([c.de(v) for c, v in zip(celdas, valores)])
        total += 1
    if columnas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{total + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def plantilla_a_csv(columnas: list[Columna], renglones: Iterable[list[Any]]) -> bytes:
    """Mismas columnas que el Excel, con BOM UTF-8; fechas dd/mm/aaaa."""
    buf = io.StringIO()
    buf.write("\ufeff")  # BOM
    writer = csv.writer(buf)
    writer.writerow([texto(c.encabezado, "texto") for c in columnas])
    for valores in renglones:
        writer.writerow([texto(v, col.campo.tipo) for v, col in zip(valores, columnas)])
    return buf.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _estilo_encabezado() -> tuple[Font, PatternFill]:
    font = Font(
        bold=True,
        color=_BRAND_PRIMARY_FOREGROUND,
        name=_BRAND_FONT_NAME,
        size=_BRAND_FONT_SIZE,
    )
    return font, PatternFill("solid", fgColor=_BRAND_PRIMARY)


def _celda(ws, valor: Any, tipo: str = "texto") -> Any:
    """Valor para `ws.append` en modo write_only. Un texto que empieza con «=»
    se fuerza a texto (openpyxl lo escribiría como fórmula); fechas e importes
    llevan su formato de número."""
    if isinstance(valor, str):
        if valor.startswith("="):
            c = WriteOnlyCell(ws, value=valor)
            c.data_type = "s"
            return c
        return valor
    formato = _FORMATO_XLSX.get(tipo)
    if formato and isinstance(valor, (int, float, datetime)) and not isinstance(valor, bool):
        c = WriteOnlyCell(ws, value=valor)
        c.number_format = formato
        return c
    return valor


class _CeldasColumna:
    """Celdas de una columna del export por plantilla. openpyxl escribe cada
    renglón al hacer `append`, así que la celda con formato se reutiliza de
    renglón en renglón (crear una por valor era lo más lento con cientos de
    miles de renglones). Los textos se recortan al máximo de una celda de
    Excel y nunca se escriben como fórmula."""

    def __init__(self, ws, tipo: str):
        self.ws = ws
        self.formato = _FORMATO_XLSX.get(tipo)
        self._con_formato: Optional[WriteOnlyCell] = None
        self._texto: Optional[WriteOnlyCell] = None

    def de(self, valor: Any) -> Any:
        if isinstance(valor, str):
            if len(valor) > MAX_TEXTO_EXCEL:
                valor = valor[: MAX_TEXTO_EXCEL - 1] + "…"
            if not valor.startswith("="):
                return valor
            if self._texto is None:
                self._texto = WriteOnlyCell(self.ws)
            self._texto.value = valor
            self._texto.data_type = "s"
            return self._texto
        if (self.formato and isinstance(valor, (int, float, datetime))
                and not isinstance(valor, bool)):
            if self._con_formato is None:
                self._con_formato = WriteOnlyCell(self.ws)
                self._con_formato.number_format = self.formato
            self._con_formato.value = valor
            return self._con_formato
        return valor


def _csv_seguro(valor: Any) -> Any:
    """Textos del CSV que Excel tomaría como fórmula llevan «'» al inicio."""
    return texto(valor, "texto") if isinstance(valor, str) else valor


def _chunks(seq: list, size: int) -> Iterable[list]:
    for i in range(0, len(seq), size):
        yield seq[i : i + size]


def _header_row(ws, values: list, font: Font, fill: PatternFill) -> list:
    """
    Construye la fila de encabezado con la marca Señal (fondo en tinta
    `#10141B`, texto blanco en negritas). El sheet en modo
    write_only requiere `WriteOnlyCell` con referencia explícita al ws.
    """
    align = Alignment(horizontal="left", vertical="center")
    cells = []
    for v in values:
        c = WriteOnlyCell(ws, value=v)
        if isinstance(v, str) and v.startswith("="):
            c.data_type = "s"
        c.font = font
        c.fill = fill
        c.alignment = align
        cells.append(c)
    return cells
