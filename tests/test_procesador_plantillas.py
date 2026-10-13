"""Plantillas de export del procesador: catálogo, motor, archivos, store y API."""

import csv
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from sat_descarga.cli import config_store
from sat_descarga.procesador import db as db_mod
from sat_descarga.procesador import parse_cfdi, plantillas_store
from sat_descarga.procesador.exportar import plantilla_a_csv, plantilla_a_xlsx, to_xlsx
from sat_descarga.procesador.indexar import indexar
from sat_descarga.procesador import plantillas as plantillas_mod
from sat_descarga.procesador.plantillas import (
    CLASICO,
    COMPLETO,
    Plantilla,
    columnas_de,
    contar_renglones,
    normalizar_columnas,
    renglones,
)

MI_RFC = "BBB020202BBB"
EMISOR = "AAA010101AAA"
CLIENTE = "CCC030303CCC"
U_I = "11111111-1111-1111-1111-111111111111"
U_P = "22222222-2222-2222-2222-222222222222"
U_REL = "99999999-9999-9999-9999-999999999999"
D1, D2, D3 = (f"DDDDDDDD-0000-0000-0000-00000000000{i}" for i in (1, 2, 3))

# Encabezados de la plantilla «CFDI ingresos 2026 plantilla.xlsx» del usuario.
COLUMNAS_AMERICO = [
    "Fecha Emisión", "Estatus SAT", "Serie", "Folio", "Nombre Emisor", "RFC Emisor",
    "Nombre Receptor", "RFC Receptor", "SubTotal", "Descuento", "IVA 16%", "IVA Ret.",
    "ISR Ret.", "Total", "Forma Pago", "Método Pago", "Moneda", "Tipo de cambio",
    "Tipo Comprobante", "UUID", "Validación Local", "Conceptos", "Fecha de pago",
    "Monto de pago", "Docs Relacionados",
]


def _ingreso(uuid=U_I, moneda="MXN", tipo_cambio=None, descripcion="Gasolina Magna"):
    tc = f' TipoCambio="{tipo_cambio}"' if tipo_cambio else ""
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
    xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
    Version="4.0" Serie="A" Folio="7" Fecha="2026-05-15T10:00:00" Sello="x"
    NoCertificado="0" Certificado="x" SubTotal="300.00" Descuento="10.00"
    Moneda="{moneda}"{tc} Total="320.40" TipoDeComprobante="I" Exportacion="01"
    MetodoPago="PUE" FormaPago="03" LugarExpedicion="40000" CondicionesDePago="Contado">
  <cfdi:CfdiRelacionados TipoRelacion="04">
    <cfdi:CfdiRelacionado UUID="{U_REL}"/>
  </cfdi:CfdiRelacionados>
  <cfdi:Emisor Rfc="{EMISOR}" Nombre="GASOLINERA SA DE CV" RegimenFiscal="601"/>
  <cfdi:Receptor Rfc="{MI_RFC}" Nombre="MI EMPRESA" DomicilioFiscalReceptor="40000"
                 RegimenFiscalReceptor="612" UsoCFDI="G03"/>
  <cfdi:Conceptos>
    <cfdi:Concepto ClaveProdServ="15101514" NoIdentificacion="MAGNA-01" Cantidad="10"
        ClaveUnidad="LTR" Unidad="Litro" Descripcion="{descripcion}" ValorUnitario="20.00"
        Importe="200.00" Descuento="10.00" ObjetoImp="02">
      <cfdi:Impuestos><cfdi:Traslados>
        <cfdi:Traslado Base="190.00" Impuesto="002" TipoFactor="Tasa" TasaOCuota="0.160000"
                       Importe="30.40"/>
      </cfdi:Traslados></cfdi:Impuestos>
      <cfdi:InformacionAduanera NumeroPedimento="21 47 3807 8003832"/>
    </cfdi:Concepto>
    <cfdi:Concepto ClaveProdServ="78101800" Cantidad="1" ClaveUnidad="E48"
        Descripcion="Servicio exento" ValorUnitario="100.00" Importe="100.00" ObjetoImp="02">
      <cfdi:Impuestos><cfdi:Traslados>
        <cfdi:Traslado Base="100.00" Impuesto="002" TipoFactor="Exento"/>
      </cfdi:Traslados></cfdi:Impuestos>
    </cfdi:Concepto>
  </cfdi:Conceptos>
  <cfdi:Impuestos TotalImpuestosTrasladados="30.40">
    <cfdi:Traslados>
      <cfdi:Traslado Base="190.00" Impuesto="002" TipoFactor="Tasa" TasaOCuota="0.160000"
                     Importe="30.40"/>
      <cfdi:Traslado Base="100.00" Impuesto="002" TipoFactor="Exento"/>
    </cfdi:Traslados>
  </cfdi:Impuestos>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{uuid}" FechaTimbrado="2026-05-15T10:05:00"
        RfcProvCertif="SAT970701NN3" SelloCFD="x" NoCertificadoSAT="0" SelloSAT="x"/>
  </cfdi:Complemento>
</cfdi:Comprobante>""".encode()


def _docto(uuid):
    return (f'<pago20:DoctoRelacionado IdDocumento="{uuid}" MonedaDR="MXN" NumParcialidad="1" '
            'ImpSaldoAnt="100" ImpPagado="100" ImpSaldoInsoluto="0" ObjetoImpDR="01"/>')


def _pago():
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
    xmlns:pago20="http://www.sat.gob.mx/Pagos20"
    xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
    Version="4.0" Serie="P" Folio="1" Fecha="2026-05-20T09:00:00" Sello="x" NoCertificado="0"
    Certificado="x" SubTotal="0" Moneda="XXX" Total="0" TipoDeComprobante="P"
    Exportacion="01" LugarExpedicion="40000">
  <cfdi:Emisor Rfc="{MI_RFC}" Nombre="MI EMPRESA" RegimenFiscal="612"/>
  <cfdi:Receptor Rfc="{CLIENTE}" Nombre="CLIENTE" DomicilioFiscalReceptor="40000"
                 RegimenFiscalReceptor="601" UsoCFDI="CP01"/>
  <cfdi:Conceptos>
    <cfdi:Concepto ClaveProdServ="84111506" Cantidad="1" ClaveUnidad="ACT" Descripcion="Pago"
                   ValorUnitario="0" Importe="0" ObjetoImp="01"/>
  </cfdi:Conceptos>
  <cfdi:Complemento>
    <pago20:Pagos Version="2.0">
      <pago20:Totales MontoTotalPagos="950.00"/>
      <pago20:Pago FechaPago="2026-05-01T12:00:00" FormaDePagoP="03" MonedaP="MXN"
                   TipoCambioP="1" Monto="100.00">{_docto(D1)}</pago20:Pago>
      <pago20:Pago FechaPago="2026-05-15T12:00:00" FormaDePagoP="03" MonedaP="USD"
                   TipoCambioP="17" Monto="50.00">{_docto(D2)}{_docto(D3)}</pago20:Pago>
    </pago20:Pagos>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{U_P}" FechaTimbrado="2026-05-20T09:01:00"
        RfcProvCertif="SAT970701NN3" SelloCFD="x" NoCertificadoSAT="0" SelloSAT="x"/>
  </cfdi:Complemento>
</cfdi:Comprobante>""".encode()


@pytest.fixture()
def db(tmp_path, monkeypatch):
    monkeypatch.setattr(db_mod, "DEFAULT_DB_PATH", tmp_path / "procesador.db")
    db_mod.resetear_singleton_para_tests()
    inst = db_mod.abrir_db()
    yield inst
    inst.close()
    db_mod.resetear_singleton_para_tests()


@pytest.fixture()
def config_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(config_store, "CONFIG_DIR", tmp_path / ".sat-descarga")


def _indexar(*xmls):
    r = indexar(MI_RFC, [(f"{i}.xml", x) for i, x in enumerate(xmls)], origen="manual")
    assert r["errores"] == [] and r["agregados"] == len(xmls), r


def _reporte(db, plantilla=CLASICO, **opciones):
    columnas = columnas_de(plantilla, **opciones)
    encabezados = [c.encabezado for c in columnas]
    filas = renglones(db, {"mi_rfc": MI_RFC}, columnas, **opciones)
    return encabezados, [dict(zip(encabezados, f)) for f in filas]


# ---------------------------------------------------------------------------
# Formato clásico (el de la app web vieja)
# ---------------------------------------------------------------------------


def test_clasico_tiene_las_25_columnas_de_la_plantilla_del_usuario():
    assert [c.encabezado for c in columnas_de(CLASICO)] == COLUMNAS_AMERICO


def test_clasico_ingreso_en_un_renglon(db):
    _indexar(_ingreso())
    _, (r,) = _reporte(db)

    assert r["Fecha Emisión"] == datetime(2026, 5, 15, 10, 0)
    assert r["Estatus SAT"] == "Por verificar"
    assert (r["Serie"], r["Folio"], r["UUID"]) == ("A", "7", U_I)
    assert (r["RFC Emisor"], r["Nombre Receptor"]) == (EMISOR, "MI EMPRESA")
    assert (r["SubTotal"], r["Descuento"], r["IVA 16%"], r["Total"]) == (300.0, 10.0, 30.4, 320.4)
    assert r["Forma Pago"] == "03 - Transferencia electrónica de fondos"
    assert r["Método Pago"] == "PUE - Pago en una sola exhibición"
    assert r["Tipo Comprobante"] == "I - Ingreso"
    assert r["Tipo de cambio"] == 1.0
    assert r["Validación Local"] == "Válido"
    assert r["Conceptos"] == "Gasolina Magna | Servicio exento"
    assert (r["Fecha de pago"], r["Monto de pago"], r["Docs Relacionados"]) == (None, None, "")


def test_rep_con_varios_pagos_suma_montos_y_lista_documentos(db):
    _indexar(_pago())
    _, (r,) = _reporte(db)
    assert r["Monto de pago"] == 150.0  # el export viejo solo tomaba el primer pago
    assert r["Fecha de pago"] == "01/05/2026 | 15/05/2026"
    assert r["Docs Relacionados"] == f"{D1} | {D2} | {D3}"

    _, (mxn,) = _reporte(db, convertir_mxn=True)
    assert mxn["Monto de pago MXN"] == 100.0 + 50.0 * 17  # cada pago, con su tipo de cambio


def test_rep_sin_documento_usa_el_primer_pago_del_parser(db):
    db.agregar([parse_cfdi(_pago())], mi_rfc=MI_RFC)  # cargado antes del índice
    _, (r,) = _reporte(db)
    assert (r["Monto de pago"], r["Docs Relacionados"]) == (100.0, D1)


def test_conceptos_repetidos_se_agrupan(db):
    xml = _ingreso().replace(b'Descripcion="Servicio exento"', b'Descripcion="Gasolina Magna"')
    _indexar(xml)
    _, (r,) = _reporte(db)
    assert r["Conceptos"] == "Gasolina Magna ×2"


# ---------------------------------------------------------------------------
# Una fila por concepto
# ---------------------------------------------------------------------------


def test_por_concepto_repite_el_cfdi_y_el_bloque_va_donde_estaba_conceptos(db):
    _indexar(_ingreso())
    encabezados, filas = _reporte(db, por_concepto=True)

    i = encabezados.index("Validación Local")
    assert encabezados[i + 1:i + 4] == ["Concepto #", "No. Identificación", "Clave SAT"]
    assert "Conceptos" not in encabezados
    assert encabezados[-3:] == ["Fecha de pago", "Monto de pago", "Docs Relacionados"]

    magna, exento = filas
    assert magna["UUID"] == exento["UUID"] == U_I and magna["Total"] == exento["Total"] == 320.4
    assert (magna["Concepto #"], exento["Concepto #"]) == ("1 de 2", "2 de 2")
    assert magna["No. Identificación"] == "MAGNA-01"
    assert (magna["Cantidad"], magna["Unidad"], magna["Valor Unitario"]) == (10.0, "Litro", 20.0)
    assert (magna["Importe concepto"], magna["Descuento concepto"]) == (200.0, 10.0)
    assert magna["Objeto Imp."] == "02 - Sí objeto de impuesto"
    assert (magna["Base IVA"], magna["Tasa IVA"], magna["IVA concepto"]) == (190.0, "16%", 30.4)
    assert (exento["Tasa IVA"], exento["IVA concepto"]) == ("Exento", None)
    assert exento["Descuento concepto"] == 0.0


def test_por_concepto_sin_documento_usa_la_tabla_de_conceptos(db):
    db.agregar([parse_cfdi(_ingreso())], mi_rfc=MI_RFC)
    _, filas = _reporte(db, por_concepto=True)
    assert [f["Descripción"] for f in filas] == ["Gasolina Magna", "Servicio exento"]
    assert filas[0]["No. Identificación"] == ""  # solo vive en el XML


def test_por_concepto_sin_bloque_deja_la_descripcion_en_conceptos(db):
    _indexar(_ingreso())
    p = Plantilla("x", "x", CLASICO.columnas, detalle_concepto=[])
    _, filas = _reporte(db, p, por_concepto=True)
    assert [f["Conceptos"] for f in filas] == ["Gasolina Magna", "Servicio exento"]


def test_plantilla_sin_conceptos_pone_el_bloque_al_final(db):
    _indexar(_ingreso())
    encabezados, _ = _reporte(db, COMPLETO, por_concepto=True)
    assert encabezados[-17] == "Concepto #" and encabezados[-1] == "ISR Ret. concepto"


# ---------------------------------------------------------------------------
# Convertir a MXN, rutas y validación
# ---------------------------------------------------------------------------


def test_convertir_a_mxn_multiplica_importes_y_marca_el_encabezado(db):
    _indexar(_ingreso(moneda="USD", tipo_cambio="17.50"))
    encabezados, (r,) = _reporte(db, convertir_mxn=True)
    assert "Total MXN" in encabezados and "Moneda" in encabezados
    assert r["Total MXN"] == pytest.approx(320.40 * 17.5)
    assert r["SubTotal MXN"] == pytest.approx(300 * 17.5)
    assert r["Tipo de cambio"] == 17.5  # el tipo de cambio no se convierte

    _, (sin,) = _reporte(db)
    assert sin["Total"] == 320.4


def test_campos_por_ruta_de_nodos(db):
    _indexar(_ingreso())
    p = Plantilla("x", "x", normalizar_columnas([
        {"campo": "ruta:Comprobante.Complemento.TimbreFiscalDigital.RfcProvCertif",
         "encabezado": "PAC"},
        {"campo": "cfdi_relacionados", "encabezado": "Relacionados"},
        {"campo": "condiciones_pago", "encabezado": ""},
    ], "cfdi"), normalizar_columnas([
        {"campo": "ruta_concepto:InformacionAduanera.NumeroPedimento", "encabezado": "Ped."},
    ], "concepto"))
    encabezados, filas = _reporte(db, p, por_concepto=True)
    assert encabezados == ["PAC", "Relacionados", "Condiciones de Pago", "Ped."]
    assert filas[0]["PAC"] == "SAT970701NN3"
    assert filas[0]["Relacionados"] == f"04 - Sustitución de los CFDI previos: {U_REL}"
    assert [f["Ped."] for f in filas] == ["21 47 3807 8003832", ""]


@pytest.mark.parametrize("columnas, nivel, mensaje", [
    ([{"campo": "no_existe"}], "cfdi", "desconocido"),
    ([{"campo": "descripcion"}], "cfdi", "no va en las columnas del CFDI"),
    ([{"campo": "uuid"}], "concepto", "no va en el detalle"),
    ([{"campo": "ruta:Comprobante..Folio"}], "cfdi", "desconocido"),
    ([{"campo": "ruta:Comprobante.Folio[0]"}], "cfdi", "desconocido"),
    ([{"campo": "uuid", "encabezado": "x" * 101}], "cfdi", "demasiado largo"),
])
def test_normalizar_columnas_rechaza(columnas, nivel, mensaje):
    with pytest.raises(ValueError, match=mensaje):
        normalizar_columnas(columnas, nivel)


def test_normalizar_columnas_encabezado_vacio_usa_la_etiqueta():
    assert normalizar_columnas([{"campo": "uuid", "encabezado": "  "}], "cfdi") == [
        {"campo": "uuid", "encabezado": "UUID"}]


# ---------------------------------------------------------------------------
# Archivos
# ---------------------------------------------------------------------------


def test_xlsx_una_hoja_con_tipos_de_excel_y_sin_formulas(db):
    _indexar(_ingreso(descripcion="=HYPERLINK(&quot;http://x&quot;,&quot;clic&quot;)"), _pago())
    columnas = columnas_de(CLASICO)
    ws = load_workbook(io.BytesIO(plantilla_a_xlsx(
        columnas, renglones(db, {"mi_rfc": MI_RFC}, columnas))))["CFDIs"]

    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:Y3"
    assert [c.value for c in ws[1]] == COLUMNAS_AMERICO
    ingreso = {h: c for h, c in zip(COLUMNAS_AMERICO, ws[2])}  # orden por fecha: el I primero
    assert ingreso["Fecha Emisión"].value == datetime(2026, 5, 15, 10, 0)
    assert ingreso["Fecha Emisión"].number_format == "dd/mm/yyyy"
    assert ingreso["Total"].value == 320.4 and ingreso["Total"].number_format == "#,##0.00"
    assert ingreso["Conceptos"].value.startswith('=HYPERLINK("http://x"')
    assert ingreso["Conceptos"].data_type == "s"  # texto del emisor, nunca fórmula


def test_csv_con_bom_fechas_y_textos_seguros(db):
    _indexar(_ingreso(descripcion="=1+1"))
    columnas = columnas_de(CLASICO)
    data = plantilla_a_csv(columnas, renglones(db, {"mi_rfc": MI_RFC}, columnas))
    assert data.startswith("﻿".encode())
    encabezados, fila = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    r = dict(zip(encabezados, fila))
    assert (r["Fecha Emisión"], r["Total"], r["Tipo de cambio"]) == ("15/05/2026", "320.40", "1.0000")
    assert r["Conceptos"] == "'=1+1 | Servicio exento"


def test_xlsx_recorta_textos_al_maximo_de_una_celda(db, monkeypatch):
    from sat_descarga.procesador import exportar

    monkeypatch.setattr(exportar, "MAX_TEXTO_EXCEL", 10)
    _indexar(_ingreso())
    columnas = columnas_de(Plantilla("x", "x", [{"campo": "conceptos"}], []))
    ws = load_workbook(io.BytesIO(plantilla_a_xlsx(
        columnas, renglones(db, {"mi_rfc": MI_RFC}, columnas))))["CFDIs"]
    assert ws["A2"].value == "Gasolina …"


def test_contar_renglones(db):
    _indexar(_ingreso(), _pago())
    assert contar_renglones(db, {"mi_rfc": MI_RFC}) == 2
    assert contar_renglones(db, {"mi_rfc": MI_RFC}, por_concepto=True) == 3
    assert contar_renglones(db, {"mi_rfc": MI_RFC, "tipo": "P"}, por_concepto=True) == 1


def test_export_de_siempre_tampoco_escribe_formulas(db):
    _indexar(_ingreso(descripcion="=1+1"))
    wb = load_workbook(io.BytesIO(to_xlsx(db, {"mi_rfc": MI_RFC})))
    descripciones = [c for c in wb["Conceptos"]["C"] if c.value == "=1+1"]
    assert descripciones and descripciones[0].data_type == "s"


# ---------------------------------------------------------------------------
# Store
# ---------------------------------------------------------------------------


def test_store_crud_y_predeterminada(config_tmp):
    inicial = plantillas_store.listar()
    assert inicial["predeterminada"] == "clasico"
    assert [p["id"] for p in inicial["plantillas"]] == ["clasico", "completo"]

    nueva = plantillas_store.crear("Inventario", [{"campo": "uuid", "encabezado": "Folio fiscal"}],
                                   [{"campo": "no_identificacion", "encabezado": "SKU"}])
    assert plantillas_store.obtener(nueva["id"]).detalle_concepto == [
        {"campo": "no_identificacion", "encabezado": "SKU"}]

    plantillas_store.fijar_predeterminada(nueva["id"])
    assert plantillas_store.predeterminada().nombre == "Inventario"

    editada = plantillas_store.actualizar(nueva["id"], "Inventario 2", [{"campo": "total"}])
    assert editada["columnas"] == [{"campo": "total", "encabezado": "Total"}]

    assert plantillas_store.eliminar(nueva["id"]) is True
    assert plantillas_store.listar()["predeterminada"] == "clasico"  # vuelve la del sistema
    assert plantillas_store.eliminar(nueva["id"]) is False


def test_store_protege_las_del_sistema_y_valida(config_tmp):
    with pytest.raises(ValueError, match="no se editan"):
        plantillas_store.actualizar("clasico", "x", [{"campo": "uuid"}])
    with pytest.raises(ValueError, match="no se borran"):
        plantillas_store.eliminar("clasico")
    with pytest.raises(ValueError, match="al menos una columna"):
        plantillas_store.crear("Vacía", [])
    with pytest.raises(ValueError, match="nombre"):
        plantillas_store.crear("  ", [{"campo": "uuid"}])
    with pytest.raises(ValueError, match="no existe"):
        plantillas_store.fijar_predeterminada("nope")


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------


@pytest.fixture()
def client(db, tmp_path, monkeypatch, config_tmp):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    from sat_descarga.api import server

    monkeypatch.setattr(config_store, "EFIRMA_DIR", tmp_path / "efirma")
    config_store.add_empresa_ciec(MI_RFC, "Mi empresa", "ciec123")
    _indexar(_ingreso(), _pago())
    yield TestClient(server.app)
    server._limpiar_session()


def test_api_campos_y_crud(client):
    campos = client.get("/procesador/cfdi/campos").json()
    assert "Complemento de pago" in campos["grupos"]
    assert {"id": "pago_monto", "nivel": "cfdi", "tipo": "importe"}.items() <= next(
        c for c in campos["campos"] if c["id"] == "pago_monto").items()

    r = client.post("/procesador/plantillas", json={
        "nombre": "Mía", "columnas": [{"campo": "uuid"}, {"campo": "total", "encabezado": "T"}]})
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    assert client.put(f"/procesador/plantillas/{pid}/predeterminada").json() == {
        "predeterminada": pid}
    assert client.put(f"/procesador/plantillas/{pid}", json={
        "nombre": "Mía", "columnas": [{"campo": "no_existe"}]}).status_code == 400
    assert client.put("/procesador/plantillas/clasico", json={
        "nombre": "x", "columnas": [{"campo": "uuid"}]}).status_code == 400
    assert client.delete(f"/procesador/plantillas/{pid}").status_code == 200
    assert client.delete(f"/procesador/plantillas/{pid}").status_code == 404


def test_api_vista_previa_con_plantilla_en_edicion(client):
    r = client.post("/procesador/cfdi/vista-previa", json={
        "rfc": MI_RFC, "columnas": [{"campo": "uuid"}, {"campo": "conceptos"}],
        "detalle_concepto": [{"campo": "descripcion"}], "por_concepto": True})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["encabezados"] == ["UUID", "Descripción"]
    # Lo más reciente primero: el REP (un concepto) y luego el ingreso (dos).
    assert body["filas"] == [[U_P, "Pago"], [U_I, "Gasolina Magna"], [U_I, "Servicio exento"]]


def test_api_exportar_con_plantilla_y_sin_ella(client):
    r = client.get("/procesador/cfdi/exportar", params={
        "rfc": MI_RFC, "formato": "xlsx", "plantilla": "clasico", "por_concepto": "true"})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["CFDIs"]
    assert wb["CFDIs"].max_row == 1 + 3  # dos conceptos del ingreso + el REP

    assert client.get("/procesador/cfdi/exportar", params={
        "rfc": MI_RFC, "plantilla": "nope"}).status_code == 404

    sin = load_workbook(io.BytesIO(client.get("/procesador/cfdi/exportar", params={
        "rfc": MI_RFC, "formato": "xlsx"}).content))
    assert sin.sheetnames == ["CFDIs", "Conceptos"]  # la API pública y la MCP no cambian


def test_api_excel_que_pasa_el_tope_de_renglones_se_rechaza(client, monkeypatch):
    monkeypatch.setattr(plantillas_mod, "MAX_FILAS_EXCEL", 2)
    r = client.get("/procesador/cfdi/exportar", params={
        "rfc": MI_RFC, "formato": "xlsx", "plantilla": "clasico", "por_concepto": "true"})
    assert r.status_code == 400 and "CSV" in r.json()["detail"]
    # En CSV no hay tope.
    assert client.get("/procesador/cfdi/exportar", params={
        "rfc": MI_RFC, "formato": "csv", "plantilla": "clasico",
        "por_concepto": "true"}).status_code == 200
